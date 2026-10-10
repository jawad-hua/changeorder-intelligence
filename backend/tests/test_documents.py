"""Offline parser, evidence, and API contract tests (no fabricated frontend data)."""
import io

import fitz
import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import ai
from app.extract import parse_document, read_document
from app.grounding import locate_quote, recover_explicit_fields, verify_field
from app.main import app
from app.schemas import FIELDS, Evidence, ExtractedField


def pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((55, 70), text)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Change Order Number")
    document.add_paragraph("CO-2026-014")
    document.add_paragraph("Approval Status")
    document.add_paragraph("Pending approval")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Schedule Impact"
    table.rows[0].cells[1].text = "3 days"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_text():
    assert read_document("order.txt", b"Change order 12") == ("[Page 1]\nChange order 12", 1)


def test_pdf_pages():
    document = parse_document("co.pdf", pdf_bytes("Change order CO-44"))
    assert document.pages == 1
    assert locate_quote(document, "CO-44", 1).page == 1
    assert locate_quote(document, "CO-44", 2) is None


def test_docx_paragraph_and_table():
    parsed = parse_document("CO.DOCX", docx_bytes())
    assert parsed.pages is None
    assert "CO-2026-014" in parsed.text
    assert "Schedule Impact | 3 days" in parsed.text
    assert locate_quote(parsed, "Pending approval").page is None
    # Word pages are not reliable; a model's guessed page is dropped while valid quotes survive.
    field, issue = verify_field("approval_status", ExtractedField(
        value="Pending approval", confidence=0.8, status="supported",
        evidence=Evidence(quote="Pending approval", page=1)), parsed)
    assert field.status == "supported" and field.evidence.page is None and issue is None


def test_docx_reject_invalid_and_legacy():
    with pytest.raises(HTTPException) as err:
        parse_document("bad.docx", b"hello")
    assert err.value.status_code == 422
    with pytest.raises(HTTPException) as err:
        parse_document("old.doc", b"hello")
    assert err.value.status_code == 415


def test_invalid_extension():
    with pytest.raises(HTTPException) as err:
        parse_document("order.exe", b"hello")
    assert err.value.status_code == 415


def test_no_text():
    with pytest.raises(HTTPException):
        parse_document("empty.txt", b" ")


def test_confidence_validation():
    with pytest.raises(ValueError):
        ExtractedField(confidence=1.4)


def test_source_normalizes_line_wrap():
    document = parse_document("order.txt", b"The client requested revised conference-room\nlayout after approval.")
    assert locate_quote(document, "client requested revised conference-room layout") is not None


def test_invented_amount_is_review():
    document = parse_document("order.txt", b"Total cost USD 3,000")
    candidate = ExtractedField(value=9000, confidence=0.95, status="supported", evidence=Evidence(quote="USD 3,000", page=1))
    field, warning = verify_field("cost_impact", candidate, document)
    assert field.status == "review"
    assert field.confidence == 0
    assert warning


def test_invented_page_is_review():
    document = parse_document("order.txt", b"Change CO-3")
    candidate = ExtractedField(value="CO-3", confidence=0.7, status="supported", evidence=Evidence(quote="CO-3", page=5))
    field, _ = verify_field("change_order_number", candidate, document)
    assert field.status == "review"


def test_labelled_recovery_from_docx():
    document = parse_document("co.docx", docx_bytes())
    fields = {key: ExtractedField(confidence=0) for key in FIELDS}
    recovered = recover_explicit_fields(document, fields)
    assert recovered["approval_status"].value == "Pending approval"
    assert recovered["approval_status"].status == "supported"


def test_labelled_recovery_from_pdf_example():
    # Check the reference fixture the user tried against the first version.
    from pathlib import Path
    path = Path(__file__).resolve().parents[3] / "Construction_Change_Order_Test.pdf"
    if not path.exists():
        pytest.skip("Optional reference fixture unavailable")
    document = parse_document("test.pdf", path.read_bytes())
    fields = {key: ExtractedField(confidence=0) for key in FIELDS}
    recovered = recover_explicit_fields(document, fields)
    assert recovered["schedule_impact_days"].value == 3
    assert recovered["approval_status"].value == "Pending approval"
    assert recovered["request_date"].value == "October 8, 2026"


def test_api_file_validation():
    with TestClient(app) as client:
        res = client.post("/api/v1/extract", files={"file": ("old.doc", b"hello")})
        assert res.status_code == 415
        res = client.post("/api/v1/extract", files={"file": ("empty.txt", b"")})
        assert res.status_code == 422


def test_api_docx_response(monkeypatch):
    # No provider network call: assert its returned fields are evidence checked.
    async def offline_provider(_client, _prompt):
        return {"fields": {
            "change_order_number": {"value": "CO-2026-014", "confidence": 0.99, "status": "supported", "evidence": {"quote": "CO-2026-014", "page": None}},
            "approval_status": {"value": "Pending approval", "confidence": 0.92, "status": "supported", "evidence": {"quote": "Pending approval", "page": None}},
        }}["fields"]

    monkeypatch.setattr(ai, "_ask_groq", offline_provider)
    monkeypatch.setattr(ai.settings, "groq_api_key", "unit-test-key")
    monkeypatch.setattr(ai.settings, "enable_recovery_pass", False)
    with TestClient(app) as client:
        res = client.post("/api/v1/extract", files={"file": ("order.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["pages"] is None
    assert data["fields"]["change_order_number"]["confidence"] == 0.95
    assert data["fields"]["approval_status"]["status"] == "supported"
    assert data["fields"]["project_name"]["status"] == "missing"
