"""Offline regression coverage for real image OCR and grounded financial recovery."""
import io
import os
import shutil
import pytest

import fitz
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app import ai
from app.extract import parse_document
from app.grounding import recover_financial_fields
from app.main import app
from app.schemas import ExtractedField, FIELDS


def _image() -> bytes:
    img = Image.new("RGB", (1400, 400), "white")
    draw = ImageDraw.Draw(img)
    fonts = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"]
    font = None
    for candidate in fonts:
        try:
            font = ImageFont.truetype(candidate, 42)
            break
        except OSError:
            continue
    if font is None:
        font = ImageFont.load_default(size=42)
    draw.text((65, 75), "Change Order Number CO-2026-014", font=font, fill="black")
    draw.text((65, 160), "Total change order cost: USD 3,000.", font=font, fill="black")
    draw.text((65, 250), "Schedule Impact: 3 days", font=font, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_uploaded_photo_is_prepared_for_groq_vision_without_local_ocr(monkeypatch):
    from app import extract
    def no_local_ocr(*_args, **_kwargs):
        raise AssertionError("Direct photo upload must not call Tesseract")
    monkeypatch.setattr(extract, "ocr_image", no_local_ocr)
    document = parse_document("site-photo.png", _image())
    assert document.kind == "image"
    assert document.text == ""  # Transcription arrives asynchronously from Groq.
    assert document.ocr_text is None
    assert len(document.visuals) == 1
    assert document.visuals[0].data_uri.startswith("data:image/jpeg;base64,")


def test_image_signature_not_extension_only():
    from fastapi import HTTPException
    import pytest
    with pytest.raises(HTTPException) as err:
        parse_document("bad.png", b"not an image")
    assert err.value.status_code == 422


@pytest.mark.skipif(not (shutil.which("tesseract") or os.path.isfile(os.environ.get("TESSERACT_CMD", ""))), reason="Install Tesseract OCR for integration tests")
def test_scan_pdf_ocr_from_raster():
    pdf = fitz.open()
    page = pdf.new_page(width=800, height=250)
    page.insert_image(fitz.Rect(30, 20, 770, 240), stream=_image())
    raw = pdf.tobytes()
    pdf.close()
    document = parse_document("scan.pdf", raw)
    assert document.pages == 1
    assert "CO-2026-014" in document.text
    assert document.visuals
    assert document.ocr_text


@pytest.mark.skipif(not (shutil.which("tesseract") or os.path.isfile(os.environ.get("TESSERACT_CMD", ""))), reason="Install Tesseract OCR for integration tests")
def test_embedded_image_in_word_ocr():
    doc = Document()
    doc.add_paragraph("Please read the attached photo below")
    doc.add_picture(io.BytesIO(_image()))
    out = io.BytesIO()
    doc.save(out)
    parsed = parse_document("order.docx", out.getvalue())
    assert "CO-2026-014" in parsed.text
    assert parsed.pages is None
    assert parsed.visuals
    assert parsed.ocr_text


def test_financial_recovery_only_explicit_total():
    text = "Cost Impact\nAdditional labor: USD 1,200. Additional materials: USD 1,800. Total change order cost: USD 3,000."
    document = parse_document("co.txt", text.encode())
    fields = {key: ExtractedField(confidence=0) for key in FIELDS}
    recovered = recover_financial_fields(document, fields)
    assert recovered["cost_impact"].value == 3000
    assert recovered["currency"].value == "USD"
    assert recovered["cost_impact"].evidence.quote.endswith("USD 3,000.")


def test_recovery_does_not_guess_from_line_items_or_conflicts():
    fields = {key: ExtractedField(confidence=0) for key in FIELDS}
    for source in (
        "Additional labor: USD 1,200. Additional materials: USD 1,800.",
        "Total cost USD 3,000. Grand total USD 4,000.",
        "Grand total: $3,000.",
    ):
        assert recover_financial_fields(parse_document("x.txt", source.encode()), fields) == {}


def test_batch_recovers_missing_financial_fields_with_stable_results(monkeypatch):
    async def provider(_client, _prompt):
        return {}
    monkeypatch.setattr(ai, "_ask_groq", provider)
    monkeypatch.setattr(ai.settings, "groq_api_key", "local-test-key")
    monkeypatch.setattr(ai.settings, "enable_recovery_pass", False)
    with TestClient(app) as client:
        result = client.post("/api/v1/extract/batch", files=[
            ("files", (f"a-{i}.txt", b"Total change order cost: USD 3,000.", "text/plain"))
            for i in range(2)
        ])
    assert result.status_code == 200
    for document in result.json()["results"]:
        assert document["fields"]["cost_impact"]["value"] == 3000
        assert document["fields"]["currency"]["status"] == "supported"

def test_image_upload_api_exposes_groq_transcript_and_visual(monkeypatch):
    from app import main
    async def provider(_client, _prompt):
        return {}
    async def vision_provider(_data):
        return "Change Order Number CO-2026-014\nTotal change order cost: USD 3,000.\nSchedule Impact: 3 days"
    monkeypatch.setattr(ai, "_ask_groq", provider)
    monkeypatch.setattr(main, "transcribe_image", vision_provider)
    monkeypatch.setattr(ai.settings, "groq_api_key", "local-test-key")
    monkeypatch.setattr(ai.settings, "enable_recovery_pass", False)
    with TestClient(app) as client:
        response = client.post("/api/v1/extract", files={"file": ("photo.png", _image(), "image/png")})
    assert response.status_code == 200
    payload = response.json()
    assert "CO-2026-014" in payload["ocr_text"]
    assert payload["visuals"][0]["source"] == "Uploaded image"
    assert payload["fields"]["cost_impact"]["value"] == 3000
    assert payload["fields"]["currency"]["status"] == "supported"
