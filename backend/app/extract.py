"""Memory-only bounded readers for PDF, DOCX, TXT and image uploads."""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from pathlib import PurePath
from zipfile import ZipFile

import fitz
from docx import Document
from fastapi import HTTPException
from PIL import Image

from .ocr import MAX_VISUALS, VisualPreview, ocr_image, prepare_image, preview

logger = logging.getLogger(__name__)
MAX_PAGES = 100
MAX_OCR_PAGES = 12
MAX_EMBEDDED_OCR = 12
MAX_CHARACTERS = 120_000
MAX_DOCX_EXPANDED_BYTES = 30 * 1024 * 1024
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None


@dataclass(frozen=True)
class ParsedDocument:
    sections: tuple[Section, ...]
    pages: int | None
    kind: str
    visuals: tuple[VisualPreview, ...] = ()
    ocr_text: str | None = None

    @property
    def text(self) -> str:
        return "\n".join(
            f"[Page {section.page}]\n{section.text}" if section.page is not None else section.text
            for section in self.sections
        )


def _validate_text(
    sections: list[Section], kind: str, pages: int | None,
    visuals: list[VisualPreview] | None = None,
    ocr_text: str | None = None,
) -> ParsedDocument:
    nonempty = [s for s in sections if s.text.strip()]
    if not nonempty:
        raise HTTPException(422, "No readable text found, including OCR. Try a clearer document or higher-resolution image")
    if sum(len(s.text) for s in nonempty) > MAX_CHARACTERS:
        raise HTTPException(413, "Document text exceeds the 120,000-character limit")
    return ParsedDocument(tuple(nonempty), pages=pages, kind=kind, visuals=tuple(visuals or []), ocr_text=(ocr_text or "")[:20000] or None)


def _image_to_sections(content: bytes, source: str, page: int | None):
    img = prepare_image(content)
    text = ocr_image(img)
    visual = preview(img, source, page)
    return text, visual


def _read_image(content: bytes) -> ParsedDocument:
    """Validate and preview photo without invoking the local OCR engine.

    The API attaches a transcript returned by Groq Vision asynchronously.
    """
    image = prepare_image(content)
    return ParsedDocument((), pages=1, kind="image", visuals=(preview(image, "Uploaded image", 1),))


def attach_image_transcript(document: ParsedDocument, transcript: str) -> ParsedDocument:
    """Expose verified-as-transcript source text to the existing grounding pipeline."""
    if document.kind != "image":
        raise ValueError("Image transcript may only be attached to image documents")
    return _validate_text([Section(transcript, 1)], "image", 1, list(document.visuals), transcript)


def _read_pdf(content: bytes) -> ParsedDocument:
    if not content.startswith(b"%PDF-"):
        raise HTTPException(422, "File does not contain a valid PDF signature")
    try:
        with fitz.open(stream=io.BytesIO(content), filetype="pdf") as doc:
            if doc.needs_pass:
                raise HTTPException(422, "Password-protected PDFs are not supported")
            if len(doc) > MAX_PAGES:
                raise HTTPException(413, "PDF exceeds the 100-page limit")
            sections: list[Section] = []
            visuals: list[VisualPreview] = []
            scanned = 0
            embedded_count = 0
            ocr_chunks: list[str] = []
            for i, page in enumerate(doc):
                page_no = i + 1
                parts: list[str] = []
                native_text = page.get_text(sort=True)
                if native_text.strip():
                    parts.append(native_text)
                if len(native_text.strip()) < 15:
                    # Scanned PDF: rasterize only pages without a usable text layer.
                    scanned += 1
                    if scanned > MAX_OCR_PAGES:
                        raise HTTPException(413, f"Scanned PDF exceeds the {MAX_OCR_PAGES}-page OCR limit")
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                    if pix.width * pix.height > 14_000_000:
                        raise HTTPException(413, "Scanned PDF page exceeds safe OCR pixel limit")
                    raster = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    image_text = ocr_image(raster)
                    if image_text:
                        parts.append(image_text)
                        ocr_chunks.append(image_text)
                    if len(visuals) < MAX_VISUALS:
                        visuals.append(preview(raster, "Scanned PDF page", page_no))
                else:
                    # OCR actual embedded illustrations separately; the surrounding PDF
                    # may have a text layer while a figure still contains useful labels.
                    seen_xrefs: set[int] = set()
                    for image_info in page.get_images(full=True):
                        if embedded_count >= MAX_EMBEDDED_OCR:
                            break
                        xref = image_info[0]
                        if xref in seen_xrefs:
                            continue
                        seen_xrefs.add(xref)
                        raw = doc.extract_image(xref).get("image", b"")
                        if not raw or len(raw) > 6_000_000:
                            continue
                        try:
                            detected, visual = _image_to_sections(raw, "Embedded PDF image", page_no)
                        except HTTPException as exc:
                            # An unreadable illustration must not discard document text.
                            logger.info("Skipped PDF illustration OCR (HTTP %s)", exc.status_code)
                            continue
                        embedded_count += 1
                        if detected:
                            parts.append("[Embedded image OCR]\n" + detected)
                            ocr_chunks.append(detected)
                        if len(visuals) < MAX_VISUALS:
                            visuals.append(visual)
                sections.append(Section("\n".join(parts), page_no))
            return _validate_text(sections, "pdf", len(doc), visuals, "\n\n".join(ocr_chunks))
    except HTTPException:
        raise
    except Exception as exc:
        logger.warning("PDF parser failed (%s)", type(exc).__name__)
        raise HTTPException(422, "PDF could not be parsed") from exc


def _read_docx(content: bytes) -> ParsedDocument:
    if not content.startswith(b"PK\x03\x04"):
        raise HTTPException(422, "File does not contain a valid DOCX signature")
    try:
        with ZipFile(io.BytesIO(content)) as zip_file:
            infos = zip_file.infolist()
            if (
                len(infos) > 1500
                or sum(i.file_size for i in infos) > MAX_DOCX_EXPANDED_BYTES
                or any(i.file_size > 10 * 1024 * 1024 for i in infos)
                or any(i.file_size > 1_000_000 and i.file_size > 200 * max(i.compress_size, 1) for i in infos)
                or "word/document.xml" not in zip_file.namelist()
                or "[Content_Types].xml" not in zip_file.namelist()
            ):
                raise HTTPException(422, "DOCX archive is invalid or exceeds safe processing limits")
        document = Document(io.BytesIO(content))
        blocks: list[str] = []
        for item in document.iter_inner_content():
            if hasattr(item, "rows"):
                for row in item.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            elif item.text.strip():
                blocks.append(item.text)
        for section in document.sections:
            for part in (section.header, section.footer):
                for paragraph in part.paragraphs:
                    if paragraph.text.strip() and paragraph.text not in blocks:
                        blocks.append(paragraph.text)

        visuals: list[VisualPreview] = []
        ocr_chunks: list[str] = []
        # Use related image parts, not zip filenames, to ignore unrelated attachments.
        image_parts = [part for part in document.part.related_parts.values()
                       if getattr(part, "content_type", "").startswith("image/")]
        for part in image_parts[:MAX_EMBEDDED_OCR]:
            if len(part.blob) > 6_000_000:
                continue
            try:
                detected, visual = _image_to_sections(part.blob, "Embedded Word image", None)
            except HTTPException as exc:
                logger.info("Skipped DOCX illustration OCR (HTTP %s)", exc.status_code)
                continue
            if detected:
                blocks.append("[Embedded image OCR]\n" + detected)
                ocr_chunks.append(detected)
            if len(visuals) < MAX_VISUALS:
                visuals.append(visual)
        # DOCX page numbers depend on the layout engine; never fabricate them.
        return _validate_text([Section("\n".join(blocks), None)], "docx", None, visuals, "\n\n".join(ocr_chunks))
    except HTTPException:
        raise
    except Exception as exc:
        logger.warning("DOCX parser failed (%s)", type(exc).__name__)
        raise HTTPException(422, "DOCX could not be parsed") from exc


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    extension = PurePath(filename.replace("\\", "/")).suffix.lower()
    if extension == ".pdf":
        return _read_pdf(content)
    if extension == ".docx":
        return _read_docx(content)
    if extension in IMAGE_EXTENSIONS:
        return _read_image(content)
    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "Text must be UTF-8 encoded") from exc
        return _validate_text([Section(text, 1)], "txt", 1)
    if extension == ".doc":
        raise HTTPException(415, "Legacy .doc is not supported. Save the file as .docx first")
    raise HTTPException(415, "Supported files: PDF, DOCX, TXT, JPG, PNG, WEBP and TIFF")


def read_document(filename: str, content: bytes) -> tuple[str, int | None]:
    """Compatibility helper retained from Phase 1."""
    document = parse_document(filename, content)
    return document.text, document.pages
